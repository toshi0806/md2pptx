#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""``--hide-steps`` で中間の段が非表示になることを固定するテスト（Issue #194）．

配布用のデッキは ``@step`` の最終段だけで作る。途中経過のページは板書の代わりで、
読み返す資料には邪魔になる。

**消すのではなく隠す**のが要点。スライドが並びに残るので、番号（``a:fld`` の
``slidenum``）は隠さなかった版と同じ数字になり、**配布用 PDF のページ番号が
講義版とずれない**。実測（5 枚のうち 2 枚を隠す）では macOS PowerPoint /
LibreOffice / ppt2pdf の 3 経路とも PDF は 3 ページ、刷られた番号は 1, 4, 5 で、
段階表示のスライドは講義版と同じ ``4`` になった。

番号は pptx に数字として入っていない（``md2pptx`` が書くのは ``<a:fld
type="slidenum">`` で、キャッシュ文字列は全スライド ``‹#›``）。**実際の数字を
計算するのは開く側・変換する側**なので、ここで固定できるのは ``show="0"`` が
立つことだけ。PDF のページ数と刷られた番号は手元の 3 経路で確かめた。

テーマは python-pptx 同梱の既定テンプレートから作るので、リポジトリの
``OfficeTheme.pptx`` にも実 PowerPoint にも依存しない。
"""
from __future__ import annotations

from pptx import Presentation

from md2pptx import render
from md2pptx.parser import parse

_FM = "---\ntheme: t.pptx\n---\n\n"


def _theme(tmp_path):
    path = tmp_path / "theme.pptx"
    Presentation().save(str(path))
    return str(path)


def _shows(tmp_path, src, hide_steps):
    """生成した pptx の各スライドの ``show`` 属性を返す（未指定は None）．"""
    out = tmp_path / "out.pptx"
    render.build(parse(src), _theme(tmp_path), str(out),
                 hide_steps=hide_steps)
    return [s.element.get("show") for s in Presentation(str(out)).slides]


_STEPS = _FM + ("## 段階表示のスライド\n\n"
                "- 1段目\n"
                "<!-- @step -->\n"
                "- 2段目\n"
                "<!-- @step -->\n"
                "- 3段目\n")


def test_every_stage_but_the_last_is_hidden(tmp_path):
    """3 段のスライドは、最終段だけが表示される．"""
    assert _shows(tmp_path, _STEPS, hide_steps=True) == ["0", "0", None]


def test_without_the_flag_nothing_is_hidden(tmp_path):
    """既定では何も隠さない——``show`` 属性すら書かない．

    ``show`` は**既定が表示**なので ``show="1"`` を書いても結果は同じだが、
    書かなければ隠さない版の pptx が今までと同じバイト列になる。
    """
    assert _shows(tmp_path, _STEPS, hide_steps=False) == [None, None, None]


def test_the_slides_stay_in_the_deck(tmp_path):
    """**枚数は変わらない**（消さずに隠すため）．

    ここがページ番号のずれを防いでいる仕組み。隠したぶん枚数が減ると、
    同じ内容のページに講義版と別の番号が付く。
    """
    assert len(_shows(tmp_path, _STEPS, hide_steps=True)) == 3


def test_a_slide_without_steps_is_never_hidden(tmp_path):
    """``@step`` を持たないスライドは、フラグを付けても表示されたまま．

    原稿の見出しは 3 つだが、真ん中が 2 段に展開されるので 4 枚になる。
    期待値の並びは、前のスライド / 1 段目（中間段）/ 2 段目（最終段）/
    後のスライド——**隠れるのは 2 つめだけ**。
    """
    src = _FM + ("## 前のスライド\n\n- a\n\n"
                 "## 段階表示のスライド\n\n"
                 "- 1段目\n"
                 "<!-- @step -->\n"
                 "- 2段目\n\n"
                 "## 後のスライド\n\n- z\n")
    assert _shows(tmp_path, src, hide_steps=True) == [None, "0", None, None]


def test_at_least_one_slide_stays_visible(tmp_path):
    """**全スライドが段の途中になることはない．**

    表示されるスライドが 1 枚も無い pptx は、どの変換器でも PDF にならない
    （macOS PowerPoint は "did not produce a (non-empty) PDF"、LibreOffice は
    "impl_store failed"）。段の最終段は必ず残るので起きないが、
    ``intermediate_step`` の意味を変えたときに踏むので固定しておく。
    """
    shows = _shows(tmp_path, _STEPS, hide_steps=True)
    assert any(v is None for v in shows)


def test_the_figure_stages_are_hidden_too(tmp_path):
    """図の中の段（``seq`` の裸の ``@step``）も同じく隠れる．

    ブロック境界ではない段（Issue #125 の ``upto``）も ``_expand_steps`` が
    段として返すので、``intermediate_step`` が立つ。原稿の行を消す方式では
    表現できなかったのがここ。
    """
    src = _FM + ("## ラダー図\n\n"
                 "```seq\n"
                 "lifelines: C, S\n"
                 "C -> S: SYN\n"
                 "@step\n"
                 "S -> C: SYN+ACK\n"
                 "```\n")
    assert _shows(tmp_path, src, hide_steps=True) == ["0", None]
