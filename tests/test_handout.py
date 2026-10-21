#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""``--handout`` で配布版に出さないスライドが非表示になることを固定するテスト．

配布用のデッキからは 2 種類のスライドを落とす（Issue #194 / #197）。

- ``@step`` の最終段より前の段。途中経過のページは板書の代わりで、
  読み返す資料には邪魔になる。
- ``<!-- @handout-skip -->`` を書いたスライド。演習の解答など、
  講義では見せて配布版では隠すもの。

**消すのではなく隠す**のが要点。スライドが並びに残るので、番号（``a:fld`` の
``slidenum``）は隠さなかった版と同じ数字になり、**配布用 PDF のページ番号が
講義版とずれない**。実測（5 枚のうち 2 枚を隠す）では macOS PowerPoint /
LibreOffice / ppt2pdf の 3 経路とも PDF は 3 ページ、刷られた番号は 1, 4, 5 で、
段階表示のスライドは講義版と同じ ``4`` になった。

番号は pptx に数字として入っていない（``md2pptx`` が書くのは ``<a:fld
type="slidenum">`` で、キャッシュ文字列は全スライド ``‹#›``）。**実際の数字を
計算するのは開く側・変換する側**なので、ここで固定できるのは ``show="0"`` が
立つことだけ。PDF のページ数と刷られた番号は手元の 3 経路で確かめた。

テーマは python-pptx 同梱の既定テンプレートから作るので、リポジトリの
``OfficeTheme.pptx`` にも実 PowerPoint にも依存しない。
"""
from __future__ import annotations

import pytest
from pptx import Presentation

from md2pptx import render
from md2pptx.parser import parse

_FM = "---\ntheme: t.pptx\n---\n\n"


def _theme(tmp_path):
    path = tmp_path / "theme.pptx"
    Presentation().save(str(path))
    return str(path)


def _shows(tmp_path, src, handout):
    """生成した pptx の各スライドの ``show`` 属性を返す（未指定は None）．"""
    out = tmp_path / "out.pptx"
    render.build(parse(src), _theme(tmp_path), str(out), handout=handout)
    return [s.element.get("show") for s in Presentation(str(out)).slides]


_STEPS = _FM + ("## 段階表示のスライド\n\n"
                "- 1段目\n"
                "<!-- @step -->\n"
                "- 2段目\n"
                "<!-- @step -->\n"
                "- 3段目\n")

_SKIP = _FM + ("## 問題\n\n- 解いてください\n\n"
               "## 解答\n"
               "<!-- @handout-skip -->\n\n"
               "- 答え\n\n"
               "## まとめ\n\n- おわり\n")


# ---------------------------------------------------------------- @step

def test_every_stage_but_the_last_is_hidden(tmp_path):
    """3 段のスライドは、最終段だけが表示される．"""
    assert _shows(tmp_path, _STEPS, handout=True) == ["0", "0", None]


def test_without_the_flag_nothing_is_hidden(tmp_path):
    """既定では何も隠さない——``show`` 属性すら書かない．

    ``show`` は**既定が表示**なので ``show="1"`` を書いても結果は同じだが、
    書かなければ隠さない版の pptx が今までと同じバイト列になる。
    """
    assert _shows(tmp_path, _STEPS, handout=False) == [None, None, None]


def test_the_slides_stay_in_the_deck(tmp_path):
    """**枚数は変わらない**（消さずに隠すため）．

    ここがページ番号のずれを防いでいる仕組み。隠したぶん枚数が減ると、
    同じ内容のページに講義版と別の番号が付く。
    """
    assert len(_shows(tmp_path, _STEPS, handout=True)) == 3


def test_a_slide_without_steps_is_never_hidden(tmp_path):
    """``@step`` を持たないスライドは、フラグを付けても表示されたまま．

    原稿の見出しは 3 つだが、真ん中が 2 段に展開されるので 4 枚になる。
    期待値の並びは、前のスライド / 1 段目（中間段）/ 2 段目（最終段）/
    後のスライド——**隠れるのは 2 つめだけ**。
    """
    src = _FM + ("## 前のスライド\n\n- a\n\n"
                 "## 段階表示のスライド\n\n"
                 "- 1段目\n"
                 "<!-- @step -->\n"
                 "- 2段目\n\n"
                 "## 後のスライド\n\n- z\n")
    assert _shows(tmp_path, src, handout=True) == [None, "0", None, None]


def test_the_figure_stages_are_hidden_too(tmp_path):
    """図の中の段（``seq`` の裸の ``@step``）も同じく隠れる．

    ブロック境界ではない段（Issue #125 の ``upto``）も ``_expand_steps`` が
    段として返すので、``intermediate_step`` が立つ。原稿の行を消す方式では
    表現できなかったのがここ。
    """
    src = _FM + ("## ラダー図\n\n"
                 "```seq\n"
                 "lifelines: C, S\n"
                 "C -> S: SYN\n"
                 "@step\n"
                 "S -> C: SYN+ACK\n"
                 "```\n")
    assert _shows(tmp_path, src, handout=True) == ["0", None]


# -------------------------------------------------------- @handout-skip

def test_a_marked_slide_is_hidden(tmp_path):
    """``<!-- @handout-skip -->`` のスライドだけが隠れ、**前後は動かない**．

    番号合わせがいちばん効いてほしいのは演習回（「p.20 の問題を解いて」）で、
    解答を削除する方式ではそこだけ番号がずれていた（cn2026-03 は
    配布版 p.16 から講義版と食い違う）。
    """
    assert _shows(tmp_path, _SKIP, handout=True) == [None, "0", None]


def test_a_marked_slide_stays_visible_in_the_lecture_version(tmp_path):
    """``--handout`` を付けなければ**何も起きない**．

    この印は「配布版では隠す」という意思表示で、それ自体は何も隠さない。
    講義では解答を見せる。
    """
    assert _shows(tmp_path, _SKIP, handout=False) == [None, None, None]


def test_the_marker_can_sit_anywhere_in_the_slide(tmp_path):
    """印はスライドのどこに書いてもよい（``@step`` と同じ扱い）．"""
    src = _FM + ("## 解答\n\n- 答え\n<!-- @handout-skip -->\n- 続き\n\n"
                 "## まとめ\n\n- おわり\n")
    assert _shows(tmp_path, src, handout=True) == ["0", None]


def test_a_marked_step_slide_hides_all_of_its_stages(tmp_path):
    """``@step`` と併記したスライドは**全段が隠れる**．

    印はスライド 1 枚に対する指定で、段ごとに違う意味を持たない。
    最終段だけ隠すと、解答が出きる前の段が配布版に残ってしまう。
    """
    src = _FM + ("## 解答\n"
                 "<!-- @handout-skip -->\n\n"
                 "- 1段目\n"
                 "<!-- @step -->\n"
                 "- 2段目\n\n"
                 "## まとめ\n\n- おわり\n")
    assert _shows(tmp_path, src, handout=True) == ["0", "0", None]


def test_a_mistyped_marker_stops_the_build():
    """印の打ち間違いは**その場で止まる**（Issue #198 が前提）．

    黙って捨てられると結果は「**解答が配布版に載る**」。ビルドは成功し、
    枚数も番号も合うので、気づくのは配った後になる。
    """
    with pytest.raises(ValueError, match="unknown directive @handout-skp"):
        parse(_FM + "## 解答\n<!-- @handout-skp -->\n\n- 答え\n")


# ------------------------------------------------- 表示が0枚になるとき

def test_at_least_one_stage_stays_visible(tmp_path):
    """段だけのデッキでは、表示が 0 枚になりえない．

    段の最終段は必ず残るので起きないが、``intermediate_step`` の意味を
    変えたときに踏むので固定しておく。
    """
    assert any(v is None for v in _shows(tmp_path, _STEPS, handout=True))


def test_hiding_every_slide_warns_but_still_saves(tmp_path, capsys):
    """全スライドを隠したら警告を出し、**pptx は成功させる**．

    表示されるスライドが 1 枚も無い pptx は、どの変換器でも PDF にならない
    （macOS PowerPoint は "did not produce a (non-empty) PDF"、LibreOffice は
    "impl_store failed"）。どのエラーを読んでも ``@handout-skip`` の
    書きすぎだとは思い当たらないので、ここで言う。
    """
    src = _FM + ("## 解答1\n<!-- @handout-skip -->\n\n- a\n\n"
                 "## 解答2\n<!-- @handout-skip -->\n\n- b\n")
    assert _shows(tmp_path, src, handout=True) == ["0", "0"]
    assert "hid every slide" in capsys.readouterr().err


def test_nothing_is_said_when_a_slide_stays_visible(tmp_path, capsys):
    """1 枚でも表示が残れば警告しない（``@handout-skip`` を普通に使った場合）．"""
    _shows(tmp_path, _SKIP, handout=True)
    assert capsys.readouterr().err == ""
